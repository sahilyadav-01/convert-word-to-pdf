"""End-to-end integration test for the multi-format conversion pipeline."""

import os
import tempfile
import openpyxl
import pytest

from core.detector import Detector
from core.processor import PipelineProcessor


def test_full_pipeline_mixed_formats():
    # Detect all 6 sample files
    sample_files = [
        "samples/01_LHS.docx",
        "samples/02_LHS.pdf",
        "samples/03_RHS.csv",
        "samples/04_RHS.fwd",
        "samples/05_LHS.txt",
        "samples/06_RHS.xlsx",
    ]

    file_infos = [Detector.inspect_file(f) for f in sample_files]
    assert all(f is not None for f in file_infos)

    # 3 LHS files, 3 RHS files
    lhs_files = [f for f in file_infos if f.side == "LHS"]
    rhs_files = [f for f in file_infos if f.side == "RHS"]
    assert len(lhs_files) == 3
    assert len(rhs_files) == 3

    processor = PipelineProcessor(prefix="D")

    with tempfile.TemporaryDirectory() as tmpdir:
        out_xlsx = os.path.join(tmpdir, "Mixed_FWD_Output.xlsx")
        result = processor.process(file_infos, output_path=out_xlsx)

        assert os.path.exists(out_xlsx)
        assert result["records_count"] == 12  # 2 records from each of 6 files

        wb = openpyxl.load_workbook(out_xlsx)
        assert "Final Data" in wb.sheetnames
        assert "Processing Summary" in wb.sheetnames
        assert "Error Log" in wb.sheetnames

        ws = wb["Final Data"]
        rows = list(ws.iter_rows(values_only=True))[1:]  # Exclude header
        assert len(rows) == 12

        # Verify LHS records (first 6 rows) are sorted ascending (increasing chainage)
        lhs_rows = rows[:6]
        for r in lhs_rows:
            assert r[0] == "LHS"

        lhs_chainages = [float(r[2]) for r in lhs_rows]
        assert lhs_chainages == sorted(lhs_chainages)
        assert lhs_chainages == [100.0, 105.0, 110.0, 115.0, 120.0, 125.0]

        # Verify RHS records (last 6 rows) are sorted descending (decreasing chainage)
        rhs_rows = rows[6:]
        for r in rhs_rows:
            assert r[0] == "RHS"

        rhs_chainages = [float(r[2]) for r in rhs_rows]
        assert rhs_chainages == sorted(rhs_chainages, reverse=True)
        assert rhs_chainages == [500.0, 495.0, 490.0, 485.0, 480.0, 475.0]
