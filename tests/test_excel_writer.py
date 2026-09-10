"""Tests for Excel file generation and structure."""

import os
import tempfile
import openpyxl
import pytest
from src.excel_writer import ExcelWriter
from src.extractor import ExtractedRecord


def test_write_workbook():
    records = [
        ExtractedRecord(side="LHS", source_file="01_LHS.docx", d_value="D100", chainage=100.0, chainage_str="100", data_values=["0.55", "0.41"]),
        ExtractedRecord(side="LHS", source_file="02_LHS.docx", d_value="D105", chainage=105.0, chainage_str="105", data_values=["0.52", "0.38"]),
        ExtractedRecord(side="RHS", source_file="03_RHS.docx", d_value="D500", chainage=500.0, chainage_str="500", data_values=["0.61", "0.44"]),
        ExtractedRecord(side="RHS", source_file="02_RHS.docx", d_value="D495", chainage=495.0, chainage_str="495", data_values=["0.59", "0.43"]),
    ]

    summary = {
        "total_files": 4,
        "lhs_files": 2,
        "rhs_files": 2,
        "unknown_files": 0,
        "total_records": 4,
        "lhs_records": 2,
        "rhs_records": 2,
        "prefix": "D",
        "output_path": "test.xlsx"
    }

    error_logs = [
        {"file": "test_warn.docx", "level": "WARNING", "message": "Test warning", "action": "Logged"}
    ]

    with tempfile.TemporaryDirectory() as tmpdir:
        out_path = os.path.join(tmpdir, "output.xlsx")
        ExcelWriter.write_workbook(out_path, records, summary, error_logs)

        assert os.path.exists(out_path)

        wb = openpyxl.load_workbook(out_path)
        sheet_names = wb.sheetnames
        assert "Final Data" in sheet_names
        assert "Processing Summary" in sheet_names
        assert "Error Log" in sheet_names

        ws_data = wb["Final Data"]
        # Header row + 4 data rows
        assert ws_data.max_row == 5
        # Verify first data row is LHS D100
        assert ws_data.cell(row=2, column=1).value == "LHS"
        assert ws_data.cell(row=2, column=2).value == "D100"
        assert ws_data.cell(row=2, column=3).value == 100.0
        # Verify 4th data row is RHS D495
        assert ws_data.cell(row=5, column=1).value == "RHS"
        assert ws_data.cell(row=5, column=2).value == "D495"
        assert ws_data.cell(row=5, column=3).value == 495.0
