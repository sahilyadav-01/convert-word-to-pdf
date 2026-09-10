"""Consolidated multi-sheet Excel generator for FWD Data Converter Pro."""

import datetime
import os
import re
from typing import Any, Dict, List
import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.data_model import FWDRecord


class ExcelWriter:
    """Creates a styled, multi-sheet Excel workbook from standardized FWDRecord items."""

    @staticmethod
    def sanitize(val: Any) -> Any:
        """Removes null bytes and illegal control characters that crash openpyxl."""
        if isinstance(val, str):
            return ILLEGAL_CHARACTERS_RE.sub("", val).strip()
        return val

    HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # Navy Blue
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    LHS_FILL = PatternFill(start_color="EDF2F8", end_color="EDF2F8", fill_type="solid")     # Light Ice Blue
    RHS_FILL = PatternFill(start_color="F2F7F2", end_color="F2F7F2", fill_type="solid")     # Light Mint
    ZEBRA_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")   # Neutral light grey

    BORDER_THIN = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    FONT_REGULAR = Font(name="Calibri", size=10, color="000000")
    FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="000000")

    @classmethod
    def write_workbook(
        cls,
        output_path: str,
        records: List[FWDRecord],
        summary_stats: Dict[str, Any],
        error_logs: List[Dict[str, str]]
    ) -> str:
        """Generates the unified 3-sheet Excel workbook."""
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb = openpyxl.Workbook()

        # Sheet 1: Final Data
        ws_data = wb.active
        ws_data.title = "Final Data"
        cls._build_data_sheet(ws_data, records)

        # Sheet 2: Processing Summary
        ws_summary = wb.create_sheet(title="Processing Summary")
        cls._build_summary_sheet(ws_summary, summary_stats)

        # Sheet 3: Error Log
        ws_log = wb.create_sheet(title="Error Log")
        cls._build_log_sheet(ws_log, error_logs)

        wb.save(output_path)
        return os.path.abspath(output_path)

    @classmethod
    def _build_data_sheet(cls, ws, records: List[FWDRecord]):
        headers = FWDRecord.get_standard_headers()

        # Add extra columns if any record has extra_data
        max_extra = 0
        for r in records:
            if r.extra_data:
                max_extra = max(max_extra, len(r.extra_data))
        for i in range(1, max_extra + 1):
            headers.append(f"Extra {i}")

        ws.append(headers)

        # Style header row
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cls.BORDER_THIN
        ws.row_dimensions[1].height = 28

        # Populate rows
        for row_idx, r in enumerate(records, start=2):
            row_vals = r.to_row()
            while len(row_vals) < len(headers):
                row_vals.append("")
            ws.append(row_vals)

            row_fill = None
            if r.side.upper() == "LHS":
                row_fill = cls.LHS_FILL if (row_idx % 2 == 0) else None
            elif r.side.upper() == "RHS":
                row_fill = cls.RHS_FILL if (row_idx % 2 == 0) else None
            elif row_idx % 2 == 0:
                row_fill = cls.ZEBRA_FILL

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = cls.FONT_REGULAR
                cell.border = cls.BORDER_THIN
                if row_fill:
                    cell.fill = row_fill

                # Alignment & number formats
                if col_idx == 1:  # Side
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if r.side.upper() == "LHS":
                        cell.font = Font(name="Calibri", size=10, bold=True, color="1565C0")
                    elif r.side.upper() == "RHS":
                        cell.font = Font(name="Calibri", size=10, bold=True, color="2E7D32")
                elif col_idx == 2:  # Station ID
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.font = cls.FONT_BOLD
                elif col_idx == 3:  # Station / Chainage
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00"
                elif 4 <= col_idx <= 11:  # Force and D0..D6
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "0.000" if col_idx > 4 else "0.0"
                elif col_idx in (12, 13):  # Temperatures
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "0.0"
                elif col_idx in (14, 15):  # Coordinates
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "0.000000"
                elif col_idx in (16, 17, 18):  # Time, Source File, Source Format
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws.row_dimensions[row_idx].height = 20

        # Freeze top row
        ws.freeze_panes = "A2"

        # Auto-filter
        if len(headers) > 0 and len(records) > 0:
            last_col_letter = get_column_letter(len(headers))
            ws.auto_filter.ref = f"A1:{last_col_letter}{len(records) + 1}"

        cls._adjust_column_widths(ws)

    @classmethod
    def _build_summary_sheet(cls, ws, stats: Dict[str, Any]):
        ws.append(["FWD DATA CONVERTER PRO - BATCH PROCESSING SUMMARY"])
        ws.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        ws.row_dimensions[1].height = 32

        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        summary_rows = [
            ("Metric / Property", "Value"),
            ("Conversion Timestamp", now_str),
            ("Total Files Processed", stats.get("total_files", 0)),
            ("LHS Files Identified", stats.get("lhs_files", 0)),
            ("RHS Files Identified", stats.get("rhs_files", 0)),
            ("Unknown Side Files", stats.get("unknown_files", 0)),
            ("Total FWD Records Extracted", stats.get("total_records", 0)),
            ("LHS Records (Ascending / Increasing)", stats.get("lhs_records", 0)),
            ("RHS Records (Descending / Decreasing)", stats.get("rhs_records", 0)),
            ("Unsorted / Unknown Records", stats.get("unknown_records", 0)),
            ("Ordering Rule Applied", "LHS (Ascending) -> RHS (Descending)"),
            ("Extraction Filter", f"Prefix: '{stats.get('prefix', 'D')}'"),
            ("Output Excel File", stats.get("output_path", ""))
        ]

        # Format breakdown
        fmt_counts = stats.get("format_counts", {})
        for fmt, cnt in sorted(fmt_counts.items()):
            summary_rows.append((f"Format: {fmt} Files", cnt))

        start_row = 3
        for r_idx, (label, val) in enumerate(summary_rows, start=start_row):
            ws.append([label, val])
            cell_lbl = ws.cell(row=r_idx, column=1)
            cell_val = ws.cell(row=r_idx, column=2)

            if r_idx == start_row:
                cell_lbl.fill = cls.HEADER_FILL
                cell_lbl.font = cls.HEADER_FONT
                cell_val.fill = cls.HEADER_FILL
                cell_val.font = cls.HEADER_FONT
                ws.row_dimensions[r_idx].height = 24
            else:
                cell_lbl.font = cls.FONT_BOLD
                cell_lbl.border = cls.BORDER_THIN
                cell_val.font = cls.FONT_REGULAR
                cell_val.border = cls.BORDER_THIN
                if (r_idx - start_row) % 2 == 0:
                    cell_lbl.fill = cls.ZEBRA_FILL
                    cell_val.fill = cls.ZEBRA_FILL
                ws.row_dimensions[r_idx].height = 20

        cls._adjust_column_widths(ws)

    @classmethod
    def _build_log_sheet(cls, ws, error_logs: List[Dict[str, str]]):
        headers = ["Log Time", "Source File", "Level", "Message", "Action Taken"]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill = cls.HEADER_FILL
            c.font = cls.HEADER_FONT
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = cls.BORDER_THIN
        ws.row_dimensions[1].height = 26

        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        if not error_logs:
            ws.append([now_str, "N/A", "INFO", "All input files processed successfully without errors.", "Completed"])
            ws.cell(row=2, column=3).font = Font(name="Calibri", size=10, color="2E7D32")
        else:
            for r_idx, item in enumerate(error_logs, start=2):
                ws.append([
                    item.get("timestamp", now_str),
                    item.get("file", ""),
                    item.get("level", "WARNING"),
                    item.get("message", ""),
                    item.get("action", "")
                ])
                for c_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=r_idx, column=c_idx)
                    cell.border = cls.BORDER_THIN
                    cell.font = cls.FONT_REGULAR
                    if c_idx == 3:
                        lvl = item.get("level", "WARNING").upper()
                        if lvl == "ERROR":
                            cell.font = Font(name="Calibri", size=10, bold=True, color="C00000")
                        else:
                            cell.font = Font(name="Calibri", size=10, bold=True, color="ED7D31")

        ws.freeze_panes = "A2"
        cls._adjust_column_widths(ws)

    @staticmethod
    def _adjust_column_widths(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = min(len(val_str), 50)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)
